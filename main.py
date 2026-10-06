import shutil
import os
import base64
import io
import re
from datetime import datetime, date, timedelta
from fastapi import FastAPI, UploadFile, File, Depends, Form, HTTPException, Header
from fastapi.responses import HTMLResponse, FileResponse
from sqlalchemy.orm import Session
from PIL import Image
from fast_alpr import ALPR
import bcrypt
import jwt

from database import engine, Base, get_db
from models import KayitliPlaka, GirisLog, Kullanici

from fastapi.responses import StreamingResponse
import csv

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

Base.metadata.create_all(bind=engine)

app = FastAPI(title="Otopark Plaka Sistemi")

alpr = ALPR(
    detector_model="yolo-v9-t-384-license-plate-end2end",
    ocr_model="cct-xs-v1-global-model",
)

# Aynı plaka bu süre içinde tekrar loglanmaz (saniye)
TEKRAR_ENGEL_SURESI = 60

# --- Güvenlik ayarları ---
GIZLI_ANAHTAR = "otopark-gizli-anahtar-degistir"  # token imzalama anahtarı
TOKEN_SURESI_SAAT = 8



# ===================== KİMLİK DOĞRULAMA YARDIMCILARI =====================

def token_uret(kullanici: Kullanici) -> str:
    """Kullanıcı için imzalı bir JWT token üretir."""
    icerik = {
        "kullanici_adi": kullanici.kullanici_adi,
        "rol": kullanici.rol,
        "exp": datetime.utcnow() + timedelta(hours=TOKEN_SURESI_SAAT),
    }
    return jwt.encode(icerik, GIZLI_ANAHTAR, algorithm="HS256")


def aktif_kullanici(authorization: str = Header(None)) -> dict:
    """İstekteki token'ı doğrular, geçerliyse kullanıcı bilgisini döner."""
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Giriş gerekli")
    token = authorization.split(" ")[1]
    try:
        icerik = jwt.decode(token, GIZLI_ANAHTAR, algorithms=["HS256"])
        return icerik  # {"kullanici_adi": ..., "rol": ...}
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Oturum süresi doldu, tekrar giriş yapın")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Geçersiz oturum")


def admin_gerekli(kullanici: dict = Depends(aktif_kullanici)) -> dict:
    """Sadece admin rolüne izin verir."""
    if kullanici["rol"] != "admin":
        raise HTTPException(status_code=403, detail="Bu işlem için yetkiniz yok")
    return kullanici


# ===================== SAYFA / SAĞLIK =====================

@app.get("/", response_class=HTMLResponse)
def ana_sayfa():
    return FileResponse("panel.html")


@app.get("/health")
def health():
    return {"status": "ok"}


# ===================== GİRİŞ (LOGIN) =====================

@app.post("/login")
def login(kullanici_adi: str = Form(...), sifre: str = Form(...), db: Session = Depends(get_db)):
    kullanici = db.query(Kullanici).filter(Kullanici.kullanici_adi == kullanici_adi).first()
    if not kullanici or not bcrypt.checkpw(sifre.encode("utf-8"), kullanici.sifre_hash.encode("utf-8")):
        raise HTTPException(status_code=401, detail="Kullanıcı adı veya şifre hatalı")

    token = token_uret(kullanici)
    return {
        "token": token,
        "kullanici_adi": kullanici.kullanici_adi,
        "rol": kullanici.rol,
    }


# ===================== PLAKA GİRİŞ KONTROLÜ =====================

@app.post("/giris")
async def giris(
    dosya: UploadFile = File(...),
    db: Session = Depends(get_db),
    kullanici: dict = Depends(aktif_kullanici),  # giriş yapılmış olmalı
):
    # 1. Fotoğrafı geçici kaydet
    gecici_yol = f"gecici_{dosya.filename}"
    with open(gecici_yol, "wb") as f:
        shutil.copyfileobj(dosya.file, f)

    # 2. Plakayı oku
    sonuclar = alpr.predict(gecici_yol)
    if not sonuclar:
        os.remove(gecici_yol)
        return {"bulundu": False, "mesaj": "Plaka okunamadı"}

    # 3. Plaka + konum
    sonuc = sonuclar[0]
    plaka = sonuc.ocr.text
    kutu = sonuc.detection.bounding_box
    plaka = plaka.replace(" ", "").upper()

    # Türk plaka formatı: 2 rakam (il) + 1-3 harf + 2-4 rakam
    turk_format = re.match(r"^(0[1-9]|[1-7][0-9]|8[01])[A-Z]{1,3}[0-9]{2,4}$", plaka)

    # Kayıtlı mı?
    kayit = db.query(KayitliPlaka).filter(KayitliPlaka.plaka == plaka).first()

    # Ne Türk formatı ne kayıtlı -> muhtemelen hatalı okuma, işleme
    if not turk_format and not kayit:
        os.remove(gecici_yol)
        return {"bulundu": False, "mesaj": "Geçerli plaka okunamadı"}

    # 4. Görselleri hazırla
    gorsel = Image.open(gecici_yol).convert("RGB")
    plaka_kirpik = gorsel.crop((kutu.x1, kutu.y1, kutu.x2, kutu.y2))

    def gorsel_to_base64(img):
        tampon = io.BytesIO()
        img.save(tampon, format="JPEG")
        return base64.b64encode(tampon.getvalue()).decode("utf-8")

    arac_foto_b64 = gorsel_to_base64(gorsel)
    plaka_foto_b64 = gorsel_to_base64(plaka_kirpik)
    os.remove(gecici_yol)

    izinli = kayit is not None

    # 5. Tekrar kontrolü: bu plaka son TEKRAR_ENGEL_SURESI saniyede loglandı mı?
    sinir = datetime.now() - timedelta(seconds=TEKRAR_ENGEL_SURESI)
    yakin_kayit = (
        db.query(GirisLog)
        .filter(GirisLog.plaka == plaka, GirisLog.zaman >= sinir)
        .first()
    )

    # Yeni kayıt sadece yakın zamanda aynı plaka yoksa açılır
    yeni_kayit_acildi = False
    if not yakin_kayit:
        log = GirisLog(plaka=plaka, izinli=izinli)
        db.add(log)
        db.commit()
        yeni_kayit_acildi = True

    # 6. Sonucu dön
    return {
        "bulundu": True,
        "yeni_kayit": yeni_kayit_acildi,
        "plaka": plaka,
        "izinli": izinli,
        "sahip": kayit.sahip_adi if kayit else None,
        "mesaj": "Kapı açıldı - giriş izni verildi" if izinli else "Yetkisiz araç - giriş reddedildi",
        "arac_foto": arac_foto_b64,
        "plaka_foto": plaka_foto_b64,
    }


# ===================== LİSTELEME (giriş yapmış herkes) =====================

@app.get("/kayitli-plakalar")
def kayitli_plakalar(db: Session = Depends(get_db), kullanici: dict = Depends(aktif_kullanici)):
    plakalar = db.query(KayitliPlaka).all()
    return [{"id": p.id, "plaka": p.plaka, "sahip": p.sahip_adi} for p in plakalar]


@app.get("/loglar")
def loglar(db: Session = Depends(get_db), kullanici: dict = Depends(aktif_kullanici)):
    kayitlar = db.query(GirisLog).order_by(GirisLog.zaman.desc()).all()
    return [
        {
            "id": k.id,
            "plaka": k.plaka,
            "zaman": k.zaman.strftime("%d.%m.%Y %H:%M:%S"),
            "izinli": k.izinli,
        }
        for k in kayitlar
    ]


@app.get("/istatistik")
def istatistik(db: Session = Depends(get_db), kullanici: dict = Depends(aktif_kullanici)):
    toplam_kayitli = db.query(KayitliPlaka).count()

    bugun = date.today()
    bugunku_loglar = db.query(GirisLog).filter(
        GirisLog.zaman >= datetime(bugun.year, bugun.month, bugun.day)
    ).all()

    bugun_toplam = len(bugunku_loglar)
    bugun_izinli = sum(1 for l in bugunku_loglar if l.izinli)
    bugun_yetkisiz = bugun_toplam - bugun_izinli

    return {
        "toplam_kayitli": toplam_kayitli,
        "bugun_toplam": bugun_toplam,
        "bugun_izinli": bugun_izinli,
        "bugun_yetkisiz": bugun_yetkisiz,
    }


@app.get("/plaka-gecmis/{plaka}")
def plaka_gecmis(plaka: str, db: Session = Depends(get_db), kullanici: dict = Depends(aktif_kullanici)):
    kayitlar = (
        db.query(GirisLog)
        .filter(GirisLog.plaka == plaka)
        .order_by(GirisLog.zaman.desc())
        .all()
    )
    return [
        {
            "zaman": k.zaman.strftime("%d.%m.%Y %H:%M:%S"),
            "izinli": k.izinli,
        }
        for k in kayitlar
    ]


# ===================== PLAKA YÖNETİMİ (sadece admin) =====================

@app.post("/plaka-ekle")
def plaka_ekle(
    plaka: str = Form(...),
    sahip: str = Form(...),
    db: Session = Depends(get_db),
    kullanici: dict = Depends(admin_gerekli),  # sadece admin
):
    plaka = plaka.replace(" ", "").upper()
    sahip = sahip.strip().title()

    mevcut = db.query(KayitliPlaka).filter(KayitliPlaka.plaka == plaka).first()
    if mevcut:
        raise HTTPException(status_code=400, detail="Bu plaka zaten kayıtlı")

    yeni = KayitliPlaka(plaka=plaka, sahip_adi=sahip)
    db.add(yeni)
    db.commit()
    return {"basarili": True, "plaka": plaka}


@app.delete("/plaka-sil/{plaka_id}")
def plaka_sil(
    plaka_id: int,
    db: Session = Depends(get_db),
    kullanici: dict = Depends(admin_gerekli),  # sadece admin
):
    kayit = db.query(KayitliPlaka).filter(KayitliPlaka.id == plaka_id).first()
    if not kayit:
        raise HTTPException(status_code=404, detail="Plaka bulunamadı")
    db.delete(kayit)
    db.commit()
    return {"basarili": True}

@app.put("/plaka-guncelle/{plaka_id}")
def plaka_guncelle(
    plaka_id: int,
    plaka: str = Form(...),
    sahip: str = Form(...),
    db: Session = Depends(get_db),
    kullanici: dict = Depends(admin_gerekli),
):
    kayit = db.query(KayitliPlaka).filter(KayitliPlaka.id == plaka_id).first()
    if not kayit:
        raise HTTPException(status_code=404, detail="Plaka bulunamadı")

    yeni_plaka = plaka.replace(" ", "").upper()

    # Başka bir kayıtta bu plaka var mı? (kendisi hariç)
    cakisma = db.query(KayitliPlaka).filter(
        KayitliPlaka.plaka == yeni_plaka, KayitliPlaka.id != plaka_id
    ).first()
    if cakisma:
        raise HTTPException(status_code=400, detail="Bu plaka zaten başka bir kayıtta var")

    kayit.plaka = yeni_plaka
    kayit.sahip_adi = sahip.strip().title()
    db.commit()
    return {"basarili": True}

@app.get("/loglar-excel")
def loglar_excel(db: Session = Depends(get_db), kullanici: dict = Depends(aktif_kullanici)):
    kayitlar = db.query(GirisLog).order_by(GirisLog.zaman.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Giriş Kayıtları"

    # Başlık satırı (koyu mavi zemin, beyaz yazı)
    basliklar = ["Plaka", "Zaman", "Durum"]
    baslik_dolgu = PatternFill(start_color="16213E", end_color="16213E", fill_type="solid")
    baslik_font = Font(color="FFFFFF", bold=True)
    for sutun, baslik in enumerate(basliklar, start=1):
        hucre = ws.cell(row=1, column=sutun, value=baslik)
        hucre.fill = baslik_dolgu
        hucre.font = baslik_font
        hucre.alignment = Alignment(horizontal="center")

    # Renk dolguları
    yesil = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    kirmizi = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")

    # Veri satırları
    for i, k in enumerate(kayitlar, start=2):
        ws.cell(row=i, column=1, value=k.plaka)
        ws.cell(row=i, column=2, value=k.zaman.strftime("%d.%m.%Y %H:%M:%S"))
        durum_hucre = ws.cell(row=i, column=3, value="İzinli" if k.izinli else "Yetkisiz")
        durum_hucre.fill = yesil if k.izinli else kirmizi

    # Sütun genişliklerini ayarla (###### sorunu olmasın)
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 12

    # Hafızada dosya oluştur
    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)

    return StreamingResponse(
        tampon,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=giris_kayitlari.xlsx"},
    )